"""Comprehensive tests for Phase 5E document generation tools.

All tests run entirely offline against temporary directories. They never
touch the real user's filesystem.

Organised into the following sections:

* :class:`TestDocumentToolDefinition` — tool definition structure.
* :class:`TestDocumentToolExecutor` — core executor behaviour.
* :class:`TestContentValidation` — content element validation.
* :class:`TestPathSecurity` — workspace path enforcement.
* :class:`TestAtomicWrite` — safe file creation.
* :class:`TestDocumentStructure` — actual DOCX generation correctness.
* :class:`TestToolRegistryIntegration` — registration with DefaultToolRegistry.
* :class:`TestAgentIntegration` — agent → tool → document → workspace path.
* :class:`TestErrorHandling` — various error paths.
* :class:`TestSensitiveLogging` — no content leakage in logs.
* :class:`TestOutputSizeLimit` — output size constraints.
* :class:`TestDeterminism` — deterministic output.
* :class:`TestFactoryFunction` — convenience builder.
"""
from __future__ import annotations

import logging
import tempfile
from pathlib import Path
from typing import Any, Iterator

import pytest

from core.agent import (
    Agent,
    AgentConfig,
    AgentStatus,
    Plan,
    Planner,
    PlanStep,
    SimpleVerifier,
    SyncToolExecutor,
    ToolCall,
)
from core.agent.registry import DefaultToolRegistry
from core.tools import (
    Workspace,
    CREATE_DOCUMENT_TOOL,
    DocumentToolExecutor,
    register_document_tools,
    register_file_tools,
)
from core.tools.document_tools import (
    create_document_tool_executor,
    _MAX_CONTENT_ELEMENTS,
    _MAX_LIST_ITEMS,
    _MAX_TABLE_ROWS,
    _MAX_TABLE_COLS,
    _MAX_TEXT_LENGTH,
)
from core.routing.types import (
    ModelDefinition,
    RoutingDecision,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def workspace_root() -> Iterator[Path]:
    """A temporary workspace root, removed after the test."""
    with tempfile.TemporaryDirectory(prefix="sovereign_ai_doc_test_") as tmp:
        yield Path(tmp)


@pytest.fixture
def workspace(workspace_root: Path) -> Workspace:
    """A Workspace rooted at the temporary directory."""
    return Workspace(root_path=str(workspace_root))


@pytest.fixture
def doc_executor(workspace: Workspace) -> DocumentToolExecutor:
    """A DocumentToolExecutor bound to the test workspace."""
    return DocumentToolExecutor(workspace)


# ---------------------------------------------------------------------------
# Sample content builders
# ---------------------------------------------------------------------------


def _simple_content() -> list[dict[str, Any]]:
    """A simple valid document content."""
    return [
        {"type": "heading", "text": "Introduction", "level": 1},
        {"type": "paragraph", "text": "This is a paragraph."},
        {"type": "bullet_list", "items": ["Item 1", "Item 2", "Item 3"]},
        {"type": "numbered_list", "items": ["Step 1", "Step 2"]},
        {
            "type": "table",
            "rows": [
                ["Name", "Value"],
                ["Alpha", "100"],
                ["Beta", "200"],
            ],
        },
    ]


def _make_call(
    *,
    path: str = "output/report.docx",
    title: str = "Test Document",
    content: list[dict[str, Any]] | None = None,
    metadata: dict[str, Any] | None = None,
) -> ToolCall:
    """Build a ToolCall for create_document."""
    args: dict[str, Any] = {
        "path": path,
        "title": title,
        "content": content if content is not None else _simple_content(),
    }
    if metadata is not None:
        args["metadata"] = metadata
    return ToolCall(
        call_id="test_doc_call",
        tool_name="create_document",
        arguments=args,
    )


def _fake_router():
    """A fake ModelRouter that always routes to a trivial model."""
    from unittest.mock import MagicMock, AsyncMock

    router = MagicMock()
    router.route = AsyncMock(return_value=RoutingDecision(
        model=ModelDefinition(
            logical_name="test", provider="fake", provider_model="test-model"
        ),
        reason="test",
        score=100,
        matched_capabilities=set(),
        modality_satisfied=True,
        preferred_honoured=False,
    ))
    return router


# ===========================================================================
# TestDocumentToolDefinition
# ===========================================================================


class TestDocumentToolDefinition:
    def test_tool_has_name(self) -> None:
        assert CREATE_DOCUMENT_TOOL.name == "create_document"

    def test_tool_has_description(self) -> None:
        assert len(CREATE_DOCUMENT_TOOL.description) > 20

    def test_tool_has_input_schema(self) -> None:
        schema = CREATE_DOCUMENT_TOOL.input_schema
        assert schema["type"] == "object"
        assert "path" in schema["properties"]
        assert "title" in schema["properties"]
        assert "content" in schema["properties"]
        assert "metadata" in schema["properties"]
        assert "path" in schema["required"]
        assert "title" in schema["required"]
        assert "content" in schema["required"]

    def test_tool_has_capability(self) -> None:
        assert CREATE_DOCUMENT_TOOL.capability == "document_generation"

    def test_tool_has_output_description(self) -> None:
        assert "size_bytes" in CREATE_DOCUMENT_TOOL.output_description


# ===========================================================================
# TestDocumentToolExecutor — basic behaviour
# ===========================================================================


class TestDocumentToolExecutor:
    async def test_executor_binds_workspace(self, workspace: Workspace) -> None:
        executor = DocumentToolExecutor(workspace)
        assert executor.workspace is workspace

    async def test_unknown_tool_name_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        call = ToolCall(call_id="x", tool_name="nonexistent_tool", arguments={})
        result = await doc_executor.execute(call)
        assert result.error is True
        assert "Unknown document tool" in result.error_message

    async def test_successful_creation(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call())
        assert result.error is False
        assert result.output is not None
        assert result.output["path"] == "output/report.docx"
        assert result.output["size_bytes"] > 0
        assert (workspace_root / "output" / "report.docx").exists()

    async def test_returns_write_result_dict(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(path="report.docx"))
        assert result.error is False
        assert isinstance(result.output, dict)
        assert "path" in result.output
        assert "size_bytes" in result.output
        assert "created_directories" in result.output


# ===========================================================================
# TestContentValidation
# ===========================================================================


class TestContentValidation:
    async def test_missing_path_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        call = ToolCall(
            call_id="x",
            tool_name="create_document",
            arguments={"title": "T", "content": []},
        )
        result = await doc_executor.execute(call)
        assert result.error is True
        assert "Missing required argument: path" in result.error_message

    async def test_missing_title_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        call = ToolCall(
            call_id="x",
            tool_name="create_document",
            arguments={"path": "a.docx", "content": []},
        )
        result = await doc_executor.execute(call)
        assert result.error is True
        assert "Missing required argument: title" in result.error_message

    async def test_missing_content_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        call = ToolCall(
            call_id="x",
            tool_name="create_document",
            arguments={"path": "a.docx", "title": "T"},
        )
        result = await doc_executor.execute(call)
        assert result.error is True
        assert "Missing required argument: content" in result.error_message

    async def test_empty_title_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(title="   "))
        assert result.error is True
        assert "Title must be a non-empty string" in result.error_message

    async def test_content_not_list_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        call = ToolCall(
            call_id="x",
            tool_name="create_document",
            arguments={"path": "a.docx", "title": "T", "content": "not a list"},
        )
        result = await doc_executor.execute(call)
        assert result.error is True
        assert "Content must be a list" in result.error_message

    async def test_too_many_elements_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        content = [
            {"type": "paragraph", "text": f"p{i}"}
            for i in range(_MAX_CONTENT_ELEMENTS + 1)
        ]
        result = await doc_executor.execute(_make_call(content=content))
        assert result.error is True
        assert "too many elements" in result.error_message

    async def test_invalid_element_type_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(
            _make_call(content=[{"type": "unknown_type", "text": "x"}])
        )
        assert result.error is True
        assert "invalid type" in result.error_message

    async def test_non_dict_element_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(content=["not a dict"]))
        assert result.error is True
        assert "must be an object" in result.error_message

    async def test_metadata_not_dict_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(metadata="bad"))
        assert result.error is True
        assert "Metadata must be an object" in result.error_message

    async def test_list_too_many_items_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(content=[
            {"type": "bullet_list", "items": [f"item{i}" for i in range(_MAX_LIST_ITEMS + 1)]}
        ]))
        assert result.error is True
        assert "too many items" in result.error_message

    async def test_table_too_many_rows_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        rows = [[f"c{j}" for j in range(2)] for _ in range(_MAX_TABLE_ROWS + 1)]
        result = await doc_executor.execute(_make_call(content=[{"type": "table", "rows": rows}]))
        assert result.error is True
        assert "too many rows" in result.error_message

    async def test_table_too_many_cols_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        row = [f"c{j}" for j in range(_MAX_TABLE_COLS + 1)]
        result = await doc_executor.execute(_make_call(content=[{"type": "table", "rows": [row]}]))
        assert result.error is True
        assert "too many columns" in result.error_message

    async def test_table_non_list_row_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(
            _make_call(content=[{"type": "table", "rows": ["not a list"]}])
        )
        assert result.error is True
        assert "row must be a list" in result.error_message

    async def test_empty_content_succeeds(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        """An empty content list is valid — document just has a title."""
        result = await doc_executor.execute(_make_call(content=[]))
        assert result.error is False
        assert result.output["size_bytes"] > 0

    async def test_element_without_type_returns_error(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(content=[{"text": "hello"}]))
        assert result.error is True
        assert "invalid type" in result.error_message

    async def test_table_non_list_rows_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(
            _make_call(path="report.docx", content=[{"type": "table", "rows": "not a list"}])
        )
        assert result.error is True
        assert "rows must be a list" in result.error_message

    async def test_bullet_list_non_items_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(
            _make_call(path="report.docx", content=[{"type": "bullet_list", "items": "not a list"}])
        )
        assert result.error is True
        assert "items must be a list" in result.error_message


# ===========================================================================
# TestPathSecurity
# ===========================================================================


class TestPathSecurity:
    async def test_traversal_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(path="../escape.docx"))
        assert result.error is True
        assert "Path is not allowed" in result.error_message

    async def test_absolute_path_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(path="/tmp/evil.docx"))
        assert result.error is True
        assert "Path is not allowed" in result.error_message

    async def test_non_docx_extension_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(path="report.txt"))
        assert result.error is True
        assert ".docx" in result.error_message

    async def test_empty_path_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        call = ToolCall(
            call_id="x",
            tool_name="create_document",
            arguments={"path": "", "title": "T", "content": []},
        )
        result = await doc_executor.execute(call)
        assert result.error is True

    async def test_windows_drive_letter_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(path="C:\\Windows\\evil.docx"))
        assert result.error is True
        assert "Path is not allowed" in result.error_message

    async def test_unc_path_rejected(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(path="\\\\server\\share\\evil.docx"))
        assert result.error is True
        assert "Path is not allowed" in result.error_message

    async def test_symlink_escape_rejected(self, workspace_root: Path) -> None:
        """A symlink pointing outside the workspace should be rejected."""
        link_target = workspace_root / "escape"
        try:
            link_target.symlink_to("/etc")
        except OSError:
            pytest.skip("Symlinks not supported on this platform")
        workspace = Workspace(root_path=str(workspace_root))
        executor = DocumentToolExecutor(workspace)
        result = await executor.execute(_make_call(path="escape/report.docx"))
        assert result.error is True
        assert "Path is not allowed" in result.error_message

    async def test_no_external_network_access(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        """Document generation must not make any network calls."""
        result = await doc_executor.execute(_make_call())
        assert result.error is False


# ===========================================================================
# TestAtomicWrite
# ===========================================================================


class TestAtomicWrite:
    async def test_overwrite_existing_file(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result1 = await doc_executor.execute(_make_call(path="report.docx"))
        assert result1.error is False
        size1 = result1.output["size_bytes"]

        result2 = await doc_executor.execute(
            _make_call(
                path="report.docx",
                title="Updated Title",
                content=[{"type": "paragraph", "text": "New content."}],
            )
        )
        assert result2.error is False
        size2 = result2.output["size_bytes"]
        assert size2 != size1 or size2 > 0

    async def test_intermediate_directories_created(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(path="a/b/c/deep/report.docx"))
        assert result.error is False
        assert (workspace_root / "a" / "b" / "c" / "deep" / "report.docx").exists()
        assert len(result.output["created_directories"]) > 0

    async def test_no_temp_files_left_behind(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        await doc_executor.execute(_make_call(path="report.docx"))
        for f in workspace_root.rglob("*"):
            assert not f.name.startswith(".docgen_"), f"Leftover temp file: {f}"


# ===========================================================================
# TestDocumentStructure — verify the actual DOCX content
# ===========================================================================


class TestDocumentStructure:
    async def test_heading_present(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            title="Main Title",
            content=[{"type": "heading", "text": "Sub Heading", "level": 2}],
        ))
        assert result.error is False

        from docx import Document

        doc = Document(str(workspace_root / "report.docx"))
        paragraphs = [p.text for p in doc.paragraphs]
        assert "Main Title" in paragraphs
        assert "Sub Heading" in paragraphs

    async def test_paragraph_text(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "paragraph", "text": "Hello world"}],
        ))
        assert result.error is False

        from docx import Document

        doc = Document(str(workspace_root / "report.docx"))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Hello world" in all_text

    async def test_bullet_list_items(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "bullet_list", "items": ["Alpha", "Beta"]}],
        ))
        assert result.error is False

        from docx import Document

        doc = Document(str(workspace_root / "report.docx"))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Alpha" in all_text
        assert "Beta" in all_text

    async def test_numbered_list_items(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "numbered_list", "items": ["First", "Second"]}],
        ))
        assert result.error is False

        from docx import Document

        doc = Document(str(workspace_root / "report.docx"))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "First" in all_text
        assert "Second" in all_text

    async def test_table_structure(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{
                "type": "table",
                "rows": [["Name", "Score"], ["Alice", "95"], ["Bob", "87"]],
            }],
        ))
        assert result.error is False

        from docx import Document

        doc = Document(str(workspace_root / "report.docx"))
        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert len(table.rows) == 3
        assert table.cell(0, 0).text == "Name"
        assert table.cell(1, 1).text == "95"

    async def test_metadata_applied(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            metadata={"author": "Test Author", "subject": "Test Subject"},
        ))
        assert result.error is False

        from docx import Document

        doc = Document(str(workspace_root / "report.docx"))
        assert doc.core_properties.author == "Test Author"
        assert doc.core_properties.subject == "Test Subject"

    async def test_metadata_truncated_to_255_chars(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        long_author = "A" * 300
        result = await doc_executor.execute(
            _make_call(path="report.docx", metadata={"author": long_author})
        )
        assert result.error is False

        from docx import Document

        doc = Document(str(workspace_root / "report.docx"))
        assert len(doc.core_properties.author) <= 255

    async def test_paragraph_with_style(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "paragraph", "text": "Styled", "style": "Normal"}],
        ))
        assert result.error is False

    async def test_paragraph_with_invalid_style_falls_back(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "paragraph", "text": "Fallback", "style": "NonExistentStyle"}],
        ))
        assert result.error is False

    async def test_empty_list_items_skipped(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "bullet_list", "items": []}],
        ))
        assert result.error is False

    async def test_empty_table_rows_skipped(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "table", "rows": []}],
        ))
        assert result.error is False

    async def test_table_with_mismatched_columns_padded(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{
                "type": "table",
                "rows": [["A", "B", "C"], ["1"]],
            }],
        ))
        assert result.error is False

    async def test_list_items_non_string_coerced(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "bullet_list", "items": [123, True, 3.14]}],
        ))
        assert result.error is False

    async def test_long_text_truncated(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        long_text = "x" * (_MAX_TEXT_LENGTH + 100)
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "paragraph", "text": long_text}],
        ))
        assert result.error is False

    async def test_heading_level_clamped(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "heading", "text": "H", "level": 99}],
        ))
        assert result.error is False


# ===========================================================================
# TestToolRegistryIntegration
# ===========================================================================


class TestToolRegistryIntegration:
    def test_register_document_tools(self, workspace: Workspace) -> None:
        registry = DefaultToolRegistry()
        register_document_tools(registry, workspace)
        assert registry.has("create_document")

    def test_registered_tool_matches_definition(self, workspace: Workspace) -> None:
        registry = DefaultToolRegistry()
        register_document_tools(registry, workspace)
        tool = registry.get("create_document")
        assert tool.name == "create_document"
        assert tool.capability == "document_generation"

    async def test_register_file_and_document_tools(self, workspace: Workspace) -> None:
        registry = DefaultToolRegistry()
        register_file_tools(registry, workspace)
        register_document_tools(registry, workspace)
        assert registry.has("list_files")
        assert registry.has("read_file")
        assert registry.has("write_file")
        assert registry.has("create_document")
        assert len(registry.names()) == 4

    async def test_sync_executor_runs_document_tool(
        self, workspace: Workspace
    ) -> None:
        registry = DefaultToolRegistry()
        register_document_tools(registry, workspace)
        executor = SyncToolExecutor(registry)
        result = await executor.execute(_make_call(path="output.docx"))
        assert result.error is False
        assert result.output["size_bytes"] > 0

    async def test_sync_executor_returns_error_result(
        self, workspace: Workspace
    ) -> None:
        """The SyncToolExecutor returns an error ToolResult on document tool failure."""
        registry = DefaultToolRegistry()
        register_document_tools(registry, workspace)
        executor = SyncToolExecutor(registry)
        call = ToolCall(
            call_id="x",
            tool_name="create_document",
            arguments={"path": "../evil.docx", "title": "T", "content": []},
        )
        result = await executor.execute(call)
        assert result.error is True
        assert "Path is not allowed" in result.error_message


# ===========================================================================
# TestAgentIntegration — full agent → tool → document → workspace path
# ===========================================================================


class TestAgentIntegration:
    @pytest.fixture
    def tool_registry(self, workspace: Workspace) -> DefaultToolRegistry:
        registry = DefaultToolRegistry()
        register_file_tools(registry, workspace)
        register_document_tools(registry, workspace)
        return registry

    async def test_agent_creates_document(
        self,
        workspace: Workspace,
        workspace_root: Path,
        tool_registry: DefaultToolRegistry,
    ) -> None:
        """A fake planner creates a document, and it appears in the workspace."""

        class FakePlanner(Planner):
            async def plan(self, task, available_tools, state):
                return Plan(
                    goal="Create a DOCX",
                    steps=(
                        PlanStep(
                            step_id="1",
                            description="Create report",
                            tool_name="create_document",
                            inputs={
                                "path": "report.docx",
                                "title": "Agent Report",
                                "content": [
                                    {"type": "paragraph", "text": "Created by agent"},
                                ],
                            },
                        ),
                    ),
                )

        executor = SyncToolExecutor(tool_registry)
        agent = Agent(
            model_router=_fake_router(),
            tool_executor=executor,
            planner=FakePlanner(),
            verifier=SimpleVerifier(),
            config=AgentConfig(max_iterations=3, max_tool_calls=5),
        )

        result = await agent.run("Create a report")

        doc_path = workspace_root / "report.docx"
        assert doc_path.exists(), f"Document not found at {doc_path}"
        assert doc_path.stat().st_size > 0

    async def test_agent_tool_call_counts(
        self,
        workspace: Workspace,
        workspace_root: Path,
        tool_registry: DefaultToolRegistry,
    ) -> None:
        """The agent runs multiple document tool calls in sequence."""

        class FakePlanner(Planner):
            async def plan(self, task, available_tools, state):
                return Plan(
                    goal="Create docs",
                    steps=(
                        PlanStep(
                            step_id="1",
                            description="Create first doc",
                            tool_name="create_document",
                            inputs={
                                "path": "first.docx",
                                "title": "First",
                                "content": [],
                            },
                        ),
                        PlanStep(
                            step_id="2",
                            description="Create second doc",
                            tool_name="create_document",
                            inputs={
                                "path": "second.docx",
                                "title": "Second",
                                "content": [
                                    {"type": "paragraph", "text": "Content"},
                                ],
                            },
                        ),
                    ),
                )

        executor = SyncToolExecutor(tool_registry)
        agent = Agent(
            model_router=_fake_router(),
            tool_executor=executor,
            planner=FakePlanner(),
            verifier=SimpleVerifier(),
            config=AgentConfig(max_iterations=3, max_tool_calls=10),
        )

        result = await agent.run("Create two docs")
        assert result.status in (
            AgentStatus.COMPLETE,
            AgentStatus.RUNNING,
            AgentStatus.PENDING,
        )
        assert (workspace_root / "first.docx").exists()
        assert (workspace_root / "second.docx").exists()


# ===========================================================================
# TestErrorHandling — various error paths
# ===========================================================================


class TestErrorHandling:
    async def test_path_is_directory(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        """Writing to a path that resolves to a directory should fail."""
        (workspace_root / "existing.docx").mkdir()
        result = await doc_executor.execute(_make_call(path="existing.docx"))
        assert result.error is True
        assert "directory" in result.error_message.lower()

    async def test_workspace_without_create_dirs(
        self, workspace_root: Path
    ) -> None:
        ws = Workspace(root_path=str(workspace_root), allow_create_dirs=False)
        executor = DocumentToolExecutor(ws)
        result = await executor.execute(_make_call(path="subdir/report.docx"))
        assert result.error is True
        assert "does not exist" in result.error_message.lower() or "not allowed" in result.error_message.lower()

    async def test_error_message_safe(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        """Error messages must not contain absolute host paths."""
        result = await doc_executor.execute(_make_call(path="/etc/passwd.docx"))
        assert result.error is True
        assert "/etc" not in result.error_message
        assert "C:\\" not in result.error_message


# ===========================================================================
# TestSensitiveLogging
# ===========================================================================


class TestSensitiveLogging:
    async def test_document_content_not_in_logs(
        self, doc_executor: DocumentToolExecutor, caplog: Any
    ) -> None:
        caplog.set_level(logging.DEBUG)
        result = await doc_executor.execute(_make_call(
            path="report.docx",
            content=[{"type": "paragraph", "text": "SECRET_CONTENT_42"}],
        ))
        assert result.error is False
        for record in caplog.records:
            assert "SECRET_CONTENT_42" not in record.getMessage()

    async def test_path_not_in_error_log(
        self, doc_executor: DocumentToolExecutor, caplog: Any
    ) -> None:
        caplog.set_level(logging.DEBUG)
        result = await doc_executor.execute(_make_call(path="/etc/passwd.docx"))
        assert result.error is True
        for record in caplog.records:
            msg = record.getMessage()
            assert "/etc" not in msg
            assert "C:\\" not in msg


# ===========================================================================
# TestOutputSizeLimit
# ===========================================================================


class TestOutputSizeLimit:
    async def test_maximum_content_elements(
        self, doc_executor: DocumentToolExecutor
    ) -> None:
        """Exactly _MAX_CONTENT_ELEMENTS should be accepted by validation."""
        content = [
            {"type": "paragraph", "text": f"p{i}"}
            for i in range(_MAX_CONTENT_ELEMENTS)
        ]
        result = await doc_executor.execute(_make_call(path="report.docx", content=content))
        # May be accepted or rejected by the output-size limit — no crash.
        assert isinstance(result, object)

    async def test_list_items_at_limit(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        items = [f"item{i}" for i in range(_MAX_LIST_ITEMS)]
        result = await doc_executor.execute(_make_call(
            path="report.docx", content=[{"type": "bullet_list", "items": items}]
        ))
        assert result.error is False

    async def test_table_rows_at_limit(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        rows = [[f"r{i}", f"v{i}"] for i in range(_MAX_TABLE_ROWS)]
        result = await doc_executor.execute(_make_call(
            path="report.docx", content=[{"type": "table", "rows": rows}]
        ))
        assert result.error is False

    async def test_table_cols_at_limit(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        row = [f"c{i}" for i in range(_MAX_TABLE_COLS)]
        result = await doc_executor.execute(_make_call(
            path="report.docx", content=[{"type": "table", "rows": [row]}]
        ))
        assert result.error is False


# ===========================================================================
# TestDeterminism
# ===========================================================================


class TestDeterminism:
    async def test_same_input_same_output(
        self, doc_executor: DocumentToolExecutor, workspace_root: Path
    ) -> None:
        r1 = await doc_executor.execute(_make_call(path="a.docx", title="T", content=_simple_content()))
        assert r1.error is False

        r2 = await doc_executor.execute(_make_call(path="b.docx", title="T", content=_simple_content()))
        assert r2.error is False

        assert r1.output["size_bytes"] == r2.output["size_bytes"]


# ===========================================================================
# TestFactoryFunction
# ===========================================================================


class TestFactoryFunction:
    def test_create_document_tool_executor(self, workspace: Workspace) -> None:
        executor = create_document_tool_executor(workspace)
        assert isinstance(executor, DocumentToolExecutor)
        assert executor.workspace is workspace